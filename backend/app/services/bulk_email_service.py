import io
import re
from pathlib import Path
from zipfile import BadZipFile

import pandas as pd
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import EmailStr, TypeAdapter, ValidationError

from app.api.schemas import BulkEmailIssue, BulkEmailResponse
from app.services.email_service import EmailService, EmailServiceError

PLACEHOLDER_PATTERN = re.compile(r"{{\s*([^{}]+?)\s*}}")


class BulkEmailService:
    def __init__(self, email_service: EmailService) -> None:
        self.email_service = email_service
        self.email_adapter = TypeAdapter(EmailStr)

    def send_bulk_email(
        self,
        filename: str,
        file_bytes: bytes,
        subject_template: str,
        body_template: str,
    ) -> BulkEmailResponse:
        self._validate_input(filename, file_bytes, subject_template, body_template)

        dataframe = self._load_dataframe(filename, file_bytes)
        if dataframe.empty:
            raise EmailServiceError(
                "Uploaded file is empty or contains no data rows.",
                status_code=400,
            )

        dataframe = self._normalize_dataframe(dataframe)
        email_column = self._find_email_column(list(dataframe.columns))
        if email_column is None:
            raise EmailServiceError(
                "The uploaded file must contain an 'email' column.",
                status_code=400,
            )

        if email_column != "email":
            dataframe = dataframe.rename(columns={email_column: "email"})

        placeholders = self._extract_placeholders(subject_template, body_template)
        self._validate_placeholders(placeholders, list(dataframe.columns))
        self.email_service.validate_settings()

        total_rows = len(dataframe.index)
        valid_rows = 0
        sent_count = 0
        failed_count = 0
        skipped_count = 0
        failures: list[BulkEmailIssue] = []
        skipped: list[BulkEmailIssue] = []

        # Row numbers are based on data rows so the first uploaded record is row 1.
        for row_number, (_, row) in enumerate(dataframe.iterrows(), start=1):
            row_data = self._normalize_row(row.to_dict())
            email = row_data.get("email", "")

            if not email:
                skipped_count += 1
                self.email_service.log_failed_email(
                    recipient="",
                    subject=subject_template,
                    error_message="Missing email.",
                )
                skipped.append(
                    self._build_issue(
                        row=row_number,
                        email="",
                        reason="Missing email.",
                    )
                )
                continue

            try:
                validated_email = str(self.email_adapter.validate_python(email))
            except ValidationError:
                skipped_count += 1
                self.email_service.log_failed_email(
                    recipient=email,
                    subject=subject_template,
                    error_message="Invalid email format.",
                )
                skipped.append(
                    self._build_issue(
                        row=row_number,
                        email=email,
                        reason="Invalid email format.",
                    )
                )
                continue

            valid_rows += 1

            try:
                subject = self._render_template(subject_template, row_data)
                body = self._render_template(body_template, row_data)
                self.email_service.send_email(validated_email, subject, body)
                sent_count += 1
            except EmailServiceError as exc:
                failed_count += 1
                failures.append(
                    self._build_issue(
                        row=row_number,
                        email=email,
                        reason=exc.message,
                    )
                )

        return BulkEmailResponse(
            success=True,
            message="Bulk email process completed.",
            total_rows=total_rows,
            valid_rows=valid_rows,
            sent_count=sent_count,
            failed_count=failed_count,
            skipped_count=skipped_count,
            failures=failures,
            skipped=skipped,
        )

    def _validate_input(
        self,
        filename: str,
        file_bytes: bytes,
        subject_template: str,
        body_template: str,
    ) -> None:
        if not (filename or "").strip():
            raise EmailServiceError("Uploaded file is required.", status_code=400)

        if not subject_template.strip():
            raise EmailServiceError(
                "subject_template cannot be empty.",
                status_code=400,
            )

        if not body_template.strip():
            raise EmailServiceError(
                "body_template cannot be empty.",
                status_code=400,
            )

        if not file_bytes:
            raise EmailServiceError("Uploaded file is empty.", status_code=400)

        extension = Path(filename).suffix.lower()
        if extension not in {".csv", ".xlsx"}:
            raise EmailServiceError(
                "Unsupported file type. Please upload a .csv or .xlsx file.",
                status_code=400,
            )

    def _load_dataframe(self, filename: str, file_bytes: bytes) -> pd.DataFrame:
        buffer = io.BytesIO(file_bytes)
        extension = Path(filename).suffix.lower()

        try:
            if extension == ".csv":
                return pd.read_csv(buffer, dtype=str, keep_default_na=False)

            return pd.read_excel(
                buffer,
                dtype=str,
                engine="openpyxl",
                keep_default_na=False,
            )
        except (
            pd.errors.EmptyDataError,
            pd.errors.ParserError,
            BadZipFile,
            InvalidFileException,
            UnicodeDecodeError,
            ValueError,
        ) as exc:
            raise EmailServiceError(
                "Unable to read the uploaded file. Ensure it is a valid CSV or XLSX file.",
                status_code=400,
            ) from exc

    def _normalize_dataframe(self, dataframe: pd.DataFrame) -> pd.DataFrame:
        cleaned = dataframe.copy()
        cleaned.columns = [
            self._normalize_column_name(column)
            for column in cleaned.columns
        ]
        return cleaned.fillna("")

    def _normalize_row(self, row: dict[str, object]) -> dict[str, str]:
        return {
            self._normalize_column_name(key): self._normalize_cell_value(value)
            for key, value in row.items()
        }

    def _normalize_column_name(self, column_name: object) -> str:
        return str(column_name).strip().lower()

    def _normalize_cell_value(self, value: object) -> str:
        if value is None:
            return ""

        return str(value).strip()

    def _find_email_column(self, columns: list[str]) -> str | None:
        for column in columns:
            if self._normalize_column_name(column) == "email":
                return column

        return None

    def _extract_placeholders(self, *templates: str) -> set[str]:
        placeholders: set[str] = set()

        for template in templates:
            matches = PLACEHOLDER_PATTERN.findall(template)
            for match in matches:
                placeholder = self._normalize_column_name(match)
                if placeholder:
                    placeholders.add(placeholder)

        return placeholders

    def _validate_placeholders(
        self,
        placeholders: set[str],
        columns: list[str],
    ) -> None:
        available_columns = {
            self._normalize_column_name(column)
            for column in columns
        }
        missing = sorted(
            placeholder
            for placeholder in placeholders
            if placeholder not in available_columns
        )

        if missing:
            raise EmailServiceError(
                "Template placeholders not found in the uploaded file columns: "
                + ", ".join(missing)
                + ".",
                status_code=400,
            )

    def _render_template(self, template: str, row_data: dict[str, str]) -> str:
        missing: set[str] = set()

        def replace(match: re.Match[str]) -> str:
            placeholder = self._normalize_column_name(match.group(1))
            if placeholder not in row_data:
                missing.add(placeholder)
                return match.group(0)

            return row_data[placeholder]

        rendered = PLACEHOLDER_PATTERN.sub(replace, template)

        if missing:
            raise EmailServiceError(
                "Template placeholders not found in row data: "
                + ", ".join(sorted(missing))
                + ".",
                status_code=400,
            )

        return rendered

    def _build_issue(self, row: int, email: str, reason: str) -> BulkEmailIssue:
        return BulkEmailIssue(
            row=row,
            email=email,
            reason=reason,
        )
